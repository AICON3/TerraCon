#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Seven merged IMU Excel workbooks -> split directly into scene-specific XLSX files

Features
----
1. Select subjects, scenes, and sensors explicitly.
2. Process only selected targets without scanning unrelated targets.
3. Overwrite existing target XLSX files.
4. Default to 0.5 s padding before/after each scene; override with --before and --after.
5. Do not re-merge, filter, interpolate, or re-align.
6. Preserve original row order and select by time window.

Examples
----
Rebuild stair_fwd pelvis for sub_025 only:
    python split_stitched_imu_by_scene.py \
      --subject 25 \
      --scene stair_fwd \
      --sensor G

Rebuild stair_fwd for all seven sensors of sub_025:
    python split_stitched_imu_by_scene.py \
      --subject 25 \
      --scene stair_fwd \
      --all-sensors

Rebuild two scenes for pelvis only in sub_025:
    python split_stitched_imu_by_scene.py \
      --subject 25 \
      --scenes stair_fwd stair_bwd \
      --sensor G

Rebuild all scenes and sensors for sub_025:
    python split_stitched_imu_by_scene.py \
      --subject 25 \
      --all-scenes \
      --all-sensors

Example: retain 1.0 s before and 2.0 s after each scene:
    python split_stitched_imu_by_scene.py \
      --subject 25 \
      --all-scenes \
      --all-sensors \
      --before 1.0 \
      --after 2.0
"""

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


BASE_DIR = Path("/mnt/newdisk/Trigait/Trigait数据库_7.20目录示范_20260721")
STITCHED_ROOT = BASE_DIR / "original_data"
TIME_ROOT = BASE_DIR / "device_export"
OUTPUT_ROOT = BASE_DIR / "original_devide"

SENSOR_MAP = {
    "G": "pelvis",
    "L1": "L_thigh",
    "L2": "L_shank",
    "L3": "L_foot",
    "R1": "R_thigh",
    "R2": "R_shank",
    "R3": "R_foot",
}

SKIP_SHEETS = {"文件信息", "未解析字节", "info", "Info", "INFO"}

OUTPUT_COLUMNS = [
    "sample_index",
    "timestamp",
    "accel_x_g",
    "accel_y_g",
    "accel_z_g",
    "gyro_x_deg_s",
    "gyro_y_deg_s",
    "gyro_z_deg_s",
    "angle_x_deg",
    "angle_y_deg",
    "angle_z_deg",
    "timestamp_device",
    "source_file",
    "time_offset_seconds",
    "time_alignment_method",
    "time_alignment_confidence",
    "alignment_anchor_count",
    "alignment_residual_seconds",
]

COLUMN_CANDIDATES = {
    "timestamp": [
        "timestamp",
        "时间戳",
        "时间戳原始字符串",
        "timestamp_aligned",
        "aligned_timestamp",
        "timestamp_device",
    ],
    "timestamp_device": [
        "timestamp_device",
        "设备时间戳",
        "时间戳原始字符串",
        "时间戳",
        "timestamp",
    ],
    "accel_x_g": ["accel_x_g", "X轴加速度(g)", "X轴加速度（g）", "加速度X(g)"],
    "accel_y_g": ["accel_y_g", "Y轴加速度(g)", "Y轴加速度（g）", "加速度Y(g)"],
    "accel_z_g": ["accel_z_g", "Z轴加速度(g)", "Z轴加速度（g）", "加速度Z(g)"],
    "gyro_x_deg_s": [
        "gyro_x_deg_s", "X轴角速度(deg/s)", "X轴角速度（deg/s）", "角速度X(deg/s)"
    ],
    "gyro_y_deg_s": [
        "gyro_y_deg_s", "Y轴角速度(deg/s)", "Y轴角速度（deg/s）", "角速度Y(deg/s)"
    ],
    "gyro_z_deg_s": [
        "gyro_z_deg_s", "Z轴角速度(deg/s)", "Z轴角速度（deg/s）", "角速度Z(deg/s)"
    ],
    "angle_x_deg": ["angle_x_deg", "X轴角度(deg)", "X轴角度（deg）", "角度X(deg)"],
    "angle_y_deg": ["angle_y_deg", "Y轴角度(deg)", "Y轴角度（deg）", "角度Y(deg)"],
    "angle_z_deg": ["angle_z_deg", "Z轴角度(deg)", "Z轴角度（deg）", "角度Z(deg)"],
    "source_file": ["source_file", "源文件", "原始文件"],
    "time_offset_seconds": ["time_offset_seconds"],
    "time_alignment_method": ["time_alignment_method"],
    "time_alignment_confidence": ["time_alignment_confidence"],
    "alignment_anchor_count": ["alignment_anchor_count"],
    "alignment_residual_seconds": ["alignment_residual_seconds"],
}


def normalize_subject(value):
    """25 / 025 / sub_025 -> sub_025"""
    text = str(value).strip()
    m = re.fullmatch(r"sub[_-]?(\d+)", text, flags=re.I)
    if m:
        return f"sub_{int(m.group(1)):03d}"
    if re.fullmatch(r"\d+", text):
        return f"sub_{int(text):03d}"
    raise ValueError(f"Cannot identify subject number: {value}")


def find_scene_time_file(subject):
    subject_num = subject.split("_")[1]
    subject_dir = TIME_ROOT / subject

    exact_candidates = [
        subject_dir / f"场景时间记录表_sub{subject_num}.xlsm",
        subject_dir / f"场景时间记录表_{subject}.xlsm",
        subject_dir / f"场景时间记录表_sub{subject_num}.xlsx",
        subject_dir / f"场景时间记录表_{subject}.xlsx",
    ]

    for p in exact_candidates:
        if p.exists():
            return p

    candidates = sorted(subject_dir.glob("场景时间记录表*.xlsm"))
    if not candidates:
        candidates = sorted(subject_dir.glob("场景时间记录表*.xlsx"))

    if not candidates:
        raise FileNotFoundError(
            f"Scene timing sheet not found: {subject_dir}/场景时间记录表*.xlsm"
        )

    return candidates[0]


def read_scene_table(filepath):
    """Read scene, start time, and end time."""
    wb = load_workbook(filepath, data_only=True, read_only=True)

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header_row_idx = None
        scene_col = start_col = end_col = None

        for i, row in enumerate(rows[:20]):
            values = [str(v).strip() if v is not None else "" for v in row]
            if "场景" in values and "开始时间" in values and "结束时间" in values:
                header_row_idx = i
                scene_col = values.index("场景")
                start_col = values.index("开始时间")
                end_col = values.index("结束时间")
                break

        if header_row_idx is None:
            continue

        scenes = []
        for row in rows[header_row_idx + 1:]:
            if len(row) <= max(scene_col, start_col, end_col):
                continue

            scene = row[scene_col]
            start = row[start_col]
            end = row[end_col]

            if scene is None or start is None or end is None:
                continue

            start_ts = pd.to_datetime(start, errors="coerce")
            end_ts = pd.to_datetime(end, errors="coerce")

            if pd.isna(start_ts) or pd.isna(end_ts):
                continue
            if end_ts <= start_ts:
                continue

            scenes.append({
                "scene": str(scene).strip(),
                "start": pd.Timestamp(start_ts),
                "end": pd.Timestamp(end_ts),
            })

        if scenes:
            return scenes

    raise ValueError(
        f"No valid scene/start/end time records found: {filepath}"
    )


def first_existing_column(df, candidates):
    for c in candidates:
        if c in df.columns:
            return c
    return None


def read_stitched_excel(filepath):
    """
    Read merged Excel workbooks.
    If row limits require multiple worksheets, concatenate them in worksheet order.
    """
    xls = pd.ExcelFile(filepath, engine="openpyxl")
    frames = []

    for sheet in xls.sheet_names:
        if sheet in SKIP_SHEETS:
            continue

        header = pd.read_excel(xls, sheet_name=sheet, nrows=0)
        ts_col = first_existing_column(header, COLUMN_CANDIDATES["timestamp"])
        if ts_col is None:
            continue

        df = pd.read_excel(xls, sheet_name=sheet)
        if not df.empty:
            frames.append(df)

    if not frames:
        raise ValueError(f"No timestamped IMU data worksheet found: {filepath}")

    return pd.concat(frames, ignore_index=True)


def build_standard_frame(df, input_excel):
    """Map merged Excel to target CSV schema without changing IMU values."""
    out = pd.DataFrame(index=df.index)

    required = [
        "timestamp",
        "accel_x_g",
        "accel_y_g",
        "accel_z_g",
        "gyro_x_deg_s",
        "gyro_y_deg_s",
        "gyro_z_deg_s",
        "angle_x_deg",
        "angle_y_deg",
        "angle_z_deg",
    ]

    source_columns = {}
    for target in required:
        source = first_existing_column(df, COLUMN_CANDIDATES[target])
        if source is None:
            raise KeyError(
                f"{input_excel.name} is missing {target} column. Available columns: {list(df.columns)}"
            )
        source_columns[target] = source

    out["timestamp"] = pd.to_datetime(
        df[source_columns["timestamp"]], errors="coerce"
    )

    for target in required[1:]:
        out[target] = df[source_columns[target]]

    ts_dev_col = first_existing_column(df, COLUMN_CANDIDATES["timestamp_device"])
    if ts_dev_col is not None:
        out["timestamp_device"] = pd.to_datetime(df[ts_dev_col], errors="coerce")
    else:
        out["timestamp_device"] = out["timestamp"]

    source_file_col = first_existing_column(df, COLUMN_CANDIDATES["source_file"])
    if source_file_col is not None:
        out["source_file"] = df[source_file_col].astype(str)
    else:
        out["source_file"] = input_excel.name

    optional_defaults = {
        "time_offset_seconds": 0.0,
        "time_alignment_method": "none_scene_cut_only",
        "time_alignment_confidence": "not_applicable",
        "alignment_anchor_count": 0,
        "alignment_residual_seconds": 0.0,
    }

    for target, default in optional_defaults.items():
        source = first_existing_column(df, COLUMN_CANDIDATES[target])
        if source is not None:
            out[target] = df[source]
        else:
            out[target] = default

    return out


def format_timestamp_milliseconds(series):
    s = pd.to_datetime(series, errors="coerce")
    result = s.dt.strftime("%Y-%m-%d %H:%M:%S.%f").str[:-3]
    return result.where(s.notna(), "")


def select_scenes(all_scenes, wanted_scenes):
    by_name = {x["scene"]: x for x in all_scenes}

    missing = [s for s in wanted_scenes if s not in by_name]
    if missing:
        raise ValueError(
            "These scenes are absent from the timing sheet: " + ", ".join(missing)
        )

    return [by_name[s] for s in wanted_scenes]


def cut_one_sensor(subject, sensor, body_name, scenes, before_padding, after_padding):
    input_excel = STITCHED_ROOT / subject / f"{subject}_{sensor}.xlsx"

    if not input_excel.exists():
        raise FileNotFoundError(f"Merged Excel workbook does not exist: {input_excel}")

    print(f"\n[{sensor} -> {body_name}] Reading {input_excel}")
    raw = read_stitched_excel(input_excel)
    std = build_standard_frame(raw, input_excel)

    invalid_ts = int(std["timestamp"].isna().sum())
    if invalid_ts:
        print(f"  Note: {invalid_ts} timestamps could not be parsed and cannot be assigned to any scene")

    before_pad = pd.Timedelta(seconds=before_padding)
    after_pad = pd.Timedelta(seconds=after_padding)

    for sc in scenes:
        scene = sc["scene"]
        start = sc["start"] - before_pad
        end = sc["end"] + after_pad

        mask = (std["timestamp"] >= start) & (std["timestamp"] < end)
        part = std.loc[mask].copy()

        # Restart numbering at 1 for each scene
        part.insert(0, "sample_index", range(1, len(part) + 1))

        part["timestamp"] = format_timestamp_milliseconds(part["timestamp"])
        part["timestamp_device"] = format_timestamp_milliseconds(
            part["timestamp_device"]
        )

        part = part[OUTPUT_COLUMNS]

        out_dir = OUTPUT_ROOT / subject / "imu" / scene
        out_dir.mkdir(parents=True, exist_ok=True)

        out_file = out_dir / f"{subject}_imu_{body_name}.xlsx"

        # Overwrite directly without checking whether the output already exists
        with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
            part.to_excel(writer, sheet_name="imu", index=False)

            ws = writer.book["imu"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            # The time column was already formatted as a full millisecond-precision string;
            # explicitly mark it as text in Excel so deleting rows and saving with Ctrl+S
            # will not auto-convert it to shortened displays such as 08:08.5.
            header_map = {cell.value: cell.column for cell in ws[1]}

            for col_name in ("timestamp", "timestamp_device"):
                if col_name in header_map:
                    col_idx = header_map[col_name]
                    for row_idx in range(2, ws.max_row + 1):
                        ws.cell(row=row_idx, column=col_idx).number_format = "@"

            from openpyxl.utils import get_column_letter

            widths = {
                "sample_index": 14,
                "timestamp": 25,
                "accel_x_g": 14,
                "accel_y_g": 14,
                "accel_z_g": 14,
                "gyro_x_deg_s": 18,
                "gyro_y_deg_s": 18,
                "gyro_z_deg_s": 18,
                "angle_x_deg": 16,
                "angle_y_deg": 16,
                "angle_z_deg": 16,
                "timestamp_device": 25,
                "source_file": 35,
                "time_offset_seconds": 20,
                "time_alignment_method": 24,
                "time_alignment_confidence": 24,
                "alignment_anchor_count": 22,
                "alignment_residual_seconds": 26,
            }

            for name, col_idx in header_map.items():
                ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(name, 16)

        print(
            f"  [Overwriting] {scene:<14} "
            f"{start.strftime('%H:%M:%S.%f')[:-3]} ~ "
            f"{end.strftime('%H:%M:%S.%f')[:-3]} | "
            f"{len(part)} rows | {out_file}"
        )


def main():
    parser = argparse.ArgumentParser(
        description="Rebuild scene XLSX from merged Excel for selected subjects/scenes/sensors."
    )

    parser.add_argument(
        "--subject",
        required=True,
        help="Subject, e.g. 25 / 025 / sub_025",
    )

    scene_group = parser.add_mutually_exclusive_group(required=True)
    scene_group.add_argument(
        "--scene",
        help="Process a single scene, e.g. stair_fwd",
    )
    scene_group.add_argument(
        "--scenes",
        nargs="+",
        help="Process multiple scenes, e.g. --scenes stair_fwd stair_bwd",
    )
    scene_group.add_argument(
        "--all-scenes",
        action="store_true",
        help="Process all scenes listed in the timing sheet",
    )

    sensor_group = parser.add_mutually_exclusive_group(required=True)
    sensor_group.add_argument(
        "--sensor",
        choices=list(SENSOR_MAP.keys()),
        help="Process only one sensor: G/L1/L2/L3/R1/R2/R3",
    )
    sensor_group.add_argument(
        "--sensors",
        nargs="+",
        choices=list(SENSOR_MAP.keys()),
        help="Process selected sensors, e.g. --sensors G L1 L2",
    )
    sensor_group.add_argument(
        "--all-sensors",
        action="store_true",
        help="Process all seven sensors",
    )

    parser.add_argument(
        "--before",
        type=float,
        default=None,
        help="Extra seconds before scene start; default 0.5",
    )
    parser.add_argument(
        "--after",
        type=float,
        default=None,
        help="Extra seconds after scene end; default 0.5",
    )
    parser.add_argument(
        "--padding",
        type=float,
        default=None,
        help="Legacy --padding option; --before/--after take precedence when provided.",
    )

    args = parser.parse_args()

    base_padding = 0.5 if args.padding is None else args.padding
    before_padding = base_padding if args.before is None else args.before
    after_padding = base_padding if args.after is None else args.after

    if args.padding is not None and args.padding < 0:
        parser.error("--padding must be non-negative")
    if before_padding < 0:
        parser.error("--before must be non-negative")
    if after_padding < 0:
        parser.error("--after must be non-negative")

    subject = normalize_subject(args.subject)

    if args.scene:
        wanted_scenes = [args.scene]
    elif args.scenes:
        wanted_scenes = args.scenes
    else:
        wanted_scenes = None

    if args.sensor:
        sensors = [args.sensor]
    elif args.sensors:
        sensors = args.sensors
    else:
        sensors = list(SENSOR_MAP.keys())

    time_file = find_scene_time_file(subject)
    all_scenes = read_scene_table(time_file)

    if args.all_scenes:
        scenes = all_scenes
        wanted_scenes = [x["scene"] for x in all_scenes]
    else:
        scenes = select_scenes(all_scenes, wanted_scenes)

    print("=" * 100)
    print("Rebuild IMU XLSX for selected scenes")
    print(f"Subject:       {subject}")
    print(f"Scenes:         {', '.join(wanted_scenes)}")
    print(f"Sensors:       {', '.join(sensors)}")
    print(f"Merged Excel:   {STITCHED_ROOT / subject}")
    print(f"Timing sheet:   {time_file}")
    print(f"Output directory:     {OUTPUT_ROOT / subject / 'imu'}")
    print(f"Time window:     before start {before_padding:.3f} s，after end {after_padding:.3f} s")
    print("Existing target XLSX files will be overwritten")
    print("=" * 100)

    for sensor in sensors:
        cut_one_sensor(
            subject=subject,
            sensor=sensor,
            body_name=SENSOR_MAP[sensor],
            scenes=scenes,
            before_padding=before_padding,
            after_padding=after_padding,
        )

    print("\nCompleted.")


if __name__ == "__main__":
    main()
